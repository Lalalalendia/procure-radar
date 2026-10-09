from __future__ import annotations

import hashlib
import json
import re
import sqlite3
from collections import Counter
from datetime import date, datetime
from itertools import islice
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


_HEADER_PATTERNS: dict[str, tuple[str, ...]] = {
    "manufacturer_inn": ("инн",),
    "manufacturer_ogrn": ("огрн",),
    "manufacturer_address": ("фактический адрес производителя",),
    "production_address": ("адрес местонахождения производственных помещений",),
    "primary_registry_number": ("первичный регистрационный номер реестровой записи",),
    "registry_number": ("реестровый номер",),
    "registry_date": ("дата внесения в реестр",),
    "valid_until": ("срок действия",),
    "actual_end_date": ("фактическая дата прекращения действия реестровой записи",),
    "okpd2_code": ("окпд2", "окпд 2"),
    "tnved_code": ("тн вэд", "тнвэд"),
    "points": ("баллы",),
    "percentage_indicator": ("процентный показатель", "процентный"),
    "conformity": ("о соответствии", "соответствие"),
}


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\xa0", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip().lower()


def _date_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    rendered = str(value).strip()
    return rendered or None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    rendered = str(value).strip()
    return rendered or None


def _looks_like_header(row: Iterable[Any]) -> int:
    values = [_norm(value) for value in row]
    score = 0
    for value in values:
        if value in {"инн", "огрн", "окпд2", "окпд 2", "тн вэд", "тнвэд"}:
            score += 1
        if "реестров" in value:
            score += 1
        if "наименование" in value:
            score += 1
        if "срок действия" in value:
            score += 1
    return score


def _find_header(ws: Any) -> tuple[int, list[Any]]:
    best: tuple[int, int, list[Any]] | None = None
    # In openpyxl read-only mode some XLSX files report ws.max_row=None.
    # Iterate lazily and cap with islice instead of relying on worksheet dimensions.
    for row_idx, row in enumerate(islice(ws.iter_rows(values_only=True), 30), 1):
        values = list(row)
        score = _looks_like_header(values)
        if best is None or score > best[0]:
            best = (score, row_idx, values)
    if best is None or best[0] < 5:
        raise ValueError("Could not identify GISP registry header row in the first 30 rows")
    return best[1], best[2]


def _header_map(headers: list[Any]) -> dict[str, int]:
    normalized = [_norm(value) for value in headers]
    result: dict[str, int] = {}

    # The official table often contains several columns whose labels include
    # "Наименование". Do not rely only on the exact word because GISP exports
    # also use labels such as "Наименование производителя", "Полное
    # наименование организации", etc.
    name_indices = [idx for idx, value in enumerate(normalized) if value == "наименование"]
    if name_indices:
        result["manufacturer_name"] = name_indices[0]
        if len(name_indices) > 1:
            result["product_name"] = name_indices[-1]

    for target, aliases in _HEADER_PATTERNS.items():
        for idx, value in enumerate(normalized):
            if target == "registry_number" and "первичный" in value:
                continue
            if any(value == alias or alias in value for alias in aliases):
                result.setdefault(target, idx)
                break

    # Manufacturer name: explicit wording used by different GISP exports.
    if "manufacturer_name" not in result:
        manufacturer_markers = (
            "производител",
            "изготовител",
            "предприят",
            "организац",
            "юридическ",
            "юридического лица",
            "ип",
        )
        for idx, value in enumerate(normalized):
            if "наименование" in value and any(marker in value for marker in manufacturer_markers):
                result["manufacturer_name"] = idx
                break

    # Product name: prefer an explicitly product-related "Наименование".
    if "product_name" not in result:
        for idx in range(len(normalized) - 1, -1, -1):
            value = normalized[idx]
            if "наименование" in value and any(
                marker in value for marker in ("продук", "товар", "издел", "промышленной продукции")
            ):
                result["product_name"] = idx
                break

    # Positional fallback for flattened/multi-level headers.
    # In the GISP registry the manufacturer's name is normally immediately
    # before its INN, while the product name is normally before OKPD2.
    def _nearest_name_before(anchor: int, *, exclude: set[int]) -> int | None:
        # First prefer any header containing "наименование" within a small
        # distance of the anchor.
        for idx in range(anchor - 1, max(-1, anchor - 8), -1):
            if idx in exclude:
                continue
            value = normalized[idx]
            if value and "наименование" in value:
                return idx
        # Some exports place only a group label or an abbreviated field name.
        # Fall back to the closest non-empty column, but avoid obvious IDs and
        # address/date/code fields.
        reject = (
            "инн", "огрн", "адрес", "номер", "дата", "срок",
            "окпд", "тн вэд", "тнвэд", "балл", "процент", "соответств",
        )
        for idx in range(anchor - 1, max(-1, anchor - 5), -1):
            if idx in exclude:
                continue
            value = normalized[idx]
            if value and not any(marker in value for marker in reject):
                return idx
        return None

    if "manufacturer_name" not in result and "manufacturer_inn" in result:
        idx = _nearest_name_before(result["manufacturer_inn"], exclude=set())
        if idx is not None:
            result["manufacturer_name"] = idx

    if "product_name" not in result and "okpd2_code" in result:
        idx = _nearest_name_before(
            result["okpd2_code"],
            exclude={result["manufacturer_name"]} if "manufacturer_name" in result else set(),
        )
        if idx is not None:
            result["product_name"] = idx

    # Last-resort distinction when the row contains multiple generic
    # "...наименование..." columns but neither was classified above.
    generic_names = [idx for idx, value in enumerate(normalized) if "наименование" in value]
    if "manufacturer_name" not in result and generic_names:
        inn_idx = result.get("manufacturer_inn")
        before_inn = [idx for idx in generic_names if inn_idx is not None and idx < inn_idx]
        result["manufacturer_name"] = before_inn[-1] if before_inn else generic_names[0]
    if "product_name" not in result and generic_names:
        candidates = [idx for idx in generic_names if idx != result.get("manufacturer_name")]
        if candidates:
            result["product_name"] = candidates[-1]

    required = {"manufacturer_name", "manufacturer_inn", "registry_number", "product_name", "okpd2_code"}
    missing = sorted(required - result.keys())
    if missing:
        detected = [f"{idx}: {value!r}" for idx, value in enumerate(normalized) if value]
        raise ValueError(
            "GISP XLSX is missing required columns after header detection: "
            + ", ".join(missing)
            + ". Detected headers: ["
            + "; ".join(detected)
            + "]"
        )
    return result


def _cell(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def import_registry_xlsx(
    conn: sqlite3.Connection,
    path: str | Path,
    *,
    scope: str = "active",
    sheet_name: str | None = None,
) -> dict[str, Any]:
    if scope not in {"active", "all"}:
        raise ValueError("scope must be 'active' or 'all'")
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(source)

    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
        header_row, headers = _find_header(ws)
        mapping = _header_map(headers)

        cur = conn.execute(
            """
            INSERT INTO gisp_import_runs(source_path, source_scope, header_json)
            VALUES(?,?,?)
            """,
            (
                str(source.resolve()),
                scope,
                json.dumps([_text(value) for value in headers], ensure_ascii=False),
            ),
        )
        run_id = int(cur.lastrowid)
        rows_seen = 0
        rows_imported = 0
        rows_skipped = 0
        skipped_reasons: Counter[str] = Counter()

        # gisp_products is the canonical one-row-per-registry projection.
        # A registry_number may have several historical rows in the XLSX, so
        # never let input order decide which row survives. Prefer the newest
        # registry_date, then the longest validity, then raw_json as a stable
        # tie-breaker. Presence fields are refreshed on every conflict so the
        # active-snapshot stale cleanup remains correct even for older rows.
        canonical_is_newer = """
            COALESCE(julianday(excluded.registry_date), -1.0)
                > COALESCE(julianday(gisp_products.registry_date), -1.0)
            OR (
                COALESCE(julianday(excluded.registry_date), -1.0)
                    = COALESCE(julianday(gisp_products.registry_date), -1.0)
                AND COALESCE(julianday(excluded.valid_until), -1.0)
                    > COALESCE(julianday(gisp_products.valid_until), -1.0)
            )
            OR (
                COALESCE(julianday(excluded.registry_date), -1.0)
                    = COALESCE(julianday(gisp_products.registry_date), -1.0)
                AND COALESCE(julianday(excluded.valid_until), -1.0)
                    = COALESCE(julianday(gisp_products.valid_until), -1.0)
                AND COALESCE(excluded.raw_json, '') > COALESCE(gisp_products.raw_json, '')
            )
        """
        canonical_fields = (
            "primary_registry_number",
            "manufacturer_name",
            "manufacturer_inn",
            "manufacturer_ogrn",
            "manufacturer_address",
            "production_address",
            "registry_date",
            "valid_until",
            "actual_end_date",
            "product_name",
            "okpd2_code",
            "tnved_code",
            "points",
            "percentage_indicator",
            "conformity",
            "raw_json",
        )
        canonical_updates = ",\n                ".join(
            f"{field}=CASE WHEN ({canonical_is_newer}) "
            f"THEN excluded.{field} ELSE gisp_products.{field} END"
            for field in canonical_fields
        )
        product_upsert_sql = f"""
            INSERT INTO gisp_products(
                registry_number, primary_registry_number,
                manufacturer_name, manufacturer_inn, manufacturer_ogrn,
                manufacturer_address, production_address,
                registry_date, valid_until, actual_end_date,
                product_name, okpd2_code, tnved_code,
                points, percentage_indicator, conformity,
                source_scope, is_active, last_seen_run_id, raw_json, updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,datetime('now'))
            ON CONFLICT(registry_number) DO UPDATE SET
                {canonical_updates},
                source_scope=excluded.source_scope,
                is_active=excluded.is_active,
                last_seen_run_id=excluded.last_seen_run_id,
                updated_at=datetime('now')
        """

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            rows_seen += 1
            registry_number = _text(_cell(row, mapping, "registry_number"))
            manufacturer_name = _text(_cell(row, mapping, "manufacturer_name"))
            product_name = _text(_cell(row, mapping, "product_name"))
            okpd2_code = _text(_cell(row, mapping, "okpd2_code"))
            if not any((registry_number, manufacturer_name, product_name, okpd2_code)):
                rows_skipped += 1
                skipped_reasons["blank"] += 1
                continue
            if not registry_number:
                rows_skipped += 1
                skipped_reasons["missing_registry_number"] += 1
                continue

            raw = {
                key: _text(_cell(row, mapping, key))
                for key in mapping
            }
            raw_json = json.dumps(raw, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
            row_key = hashlib.sha256(raw_json.encode("utf-8")).hexdigest()
            conn.execute(
                product_upsert_sql,
                (
                    registry_number,
                    _text(_cell(row, mapping, "primary_registry_number")),
                    manufacturer_name,
                    _text(_cell(row, mapping, "manufacturer_inn")),
                    _text(_cell(row, mapping, "manufacturer_ogrn")),
                    _text(_cell(row, mapping, "manufacturer_address")),
                    _text(_cell(row, mapping, "production_address")),
                    _date_text(_cell(row, mapping, "registry_date")),
                    _date_text(_cell(row, mapping, "valid_until")),
                    _date_text(_cell(row, mapping, "actual_end_date")),
                    product_name,
                    okpd2_code,
                    _text(_cell(row, mapping, "tnved_code")),
                    _text(_cell(row, mapping, "points")),
                    _text(_cell(row, mapping, "percentage_indicator")),
                    _text(_cell(row, mapping, "conformity")),
                    scope,
                    1 if scope == "active" else None,
                    run_id,
                    raw_json,
                ),
            )
            conn.execute(
                """
                INSERT INTO gisp_registry_rows(
                    row_key, registry_number, product_name, okpd2_code, tnved_code,
                    source_scope, last_seen_run_id, raw_json, updated_at
                ) VALUES(?,?,?,?,?,?,?,?,datetime('now'))
                ON CONFLICT(row_key) DO UPDATE SET
                    registry_number=excluded.registry_number,
                    product_name=excluded.product_name,
                    okpd2_code=excluded.okpd2_code,
                    tnved_code=excluded.tnved_code,
                    source_scope=excluded.source_scope,
                    last_seen_run_id=excluded.last_seen_run_id,
                    raw_json=excluded.raw_json,
                    updated_at=datetime('now')
                """,
                (
                    row_key, registry_number, product_name, okpd2_code,
                    _text(_cell(row, mapping, "tnved_code")), scope, run_id, raw_json,
                ),
            )
            rows_imported += 1

        # An active XLSX is a full snapshot. Only after a successful read do we remove
        # records that disappeared from the new active snapshot.
        stale_deleted = 0
        stale_rows_deleted = 0
        if scope == "active":
            stale_rows_deleted = conn.execute(
                """
                DELETE FROM gisp_registry_rows
                WHERE source_scope='active' AND last_seen_run_id<>?
                """,
                (run_id,),
            ).rowcount
            stale_deleted = conn.execute(
                """
                DELETE FROM gisp_products
                WHERE source_scope='active' AND last_seen_run_id<>?
                """,
                (run_id,),
            ).rowcount

        conn.execute(
            """
            UPDATE gisp_import_runs
            SET completed_at=datetime('now'), rows_seen=?, rows_imported=?, rows_skipped=?
            WHERE id=?
            """,
            (rows_seen, rows_imported, rows_skipped, run_id),
        )
        conn.commit()
        return {
            "run_id": run_id,
            "scope": scope,
            "sheet": ws.title,
            "header_row": header_row,
            "rows_seen": rows_seen,
            "rows_imported": rows_imported,
            "rows_skipped": rows_skipped,
            "stale_deleted": stale_deleted,
            "stale_rows_deleted": stale_rows_deleted,
            "skipped_reasons": dict(skipped_reasons),
            "columns": sorted(mapping),
        }
    except Exception:
        conn.rollback()
        raise
    finally:
        wb.close()


def registry_stats(conn: sqlite3.Connection) -> dict[str, Any]:
    products = int(conn.execute("SELECT COUNT(*) FROM gisp_products").fetchone()[0])
    registry_rows = int(conn.execute("SELECT COUNT(*) FROM gisp_registry_rows").fetchone()[0])
    active = int(conn.execute("SELECT COUNT(*) FROM gisp_products WHERE is_active=1").fetchone()[0])
    manufacturers = int(
        conn.execute(
            "SELECT COUNT(DISTINCT manufacturer_inn) FROM gisp_products WHERE NULLIF(manufacturer_inn,'') IS NOT NULL"
        ).fetchone()[0]
    )
    if registry_rows:
        okpd2 = int(
            conn.execute(
                "SELECT COUNT(DISTINCT okpd2_code) FROM gisp_registry_rows WHERE NULLIF(okpd2_code,'') IS NOT NULL"
            ).fetchone()[0]
        )
    else:
        okpd2 = int(
            conn.execute(
                "SELECT COUNT(DISTINCT okpd2_code) FROM gisp_products WHERE NULLIF(okpd2_code,'') IS NOT NULL"
            ).fetchone()[0]
        )
    latest = conn.execute(
        """
        SELECT id, source_path, source_scope, started_at, completed_at,
               rows_seen, rows_imported, rows_skipped
        FROM gisp_import_runs
        WHERE completed_at IS NOT NULL
        ORDER BY id DESC LIMIT 1
        """
    ).fetchone()
    return {
        "products": products,
        "registry_rows": registry_rows,
        "active_products": active,
        "manufacturers": manufacturers,
        "okpd2_codes": okpd2,
        "latest_import": dict(latest) if latest is not None else None,
    }


def okpd2_stats(
    conn: sqlite3.Connection,
    code: str,
    *,
    active_only: bool = True,
    limit: int = 20,
) -> dict[str, Any]:
    where_active = "AND p.is_active=1" if active_only else ""
    has_rows = int(conn.execute("SELECT COUNT(*) FROM gisp_registry_rows").fetchone()[0]) > 0
    if has_rows:
        summary = conn.execute(
            f"""
            SELECT COUNT(DISTINCT r.registry_number) AS products,
                   COUNT(DISTINCT NULLIF(p.manufacturer_inn,'')) AS manufacturers
            FROM gisp_registry_rows r
            JOIN gisp_products p ON p.registry_number=r.registry_number
            WHERE r.okpd2_code=? {where_active}
            """,
            (code,),
        ).fetchone()
        rows = conn.execute(
            f"""
            SELECT DISTINCT r.registry_number, r.product_name, p.manufacturer_name, p.manufacturer_inn,
                   p.valid_until, r.tnved_code
            FROM gisp_registry_rows r
            JOIN gisp_products p ON p.registry_number=r.registry_number
            WHERE r.okpd2_code=? {where_active}
            ORDER BY p.manufacturer_name, r.product_name, r.registry_number
            LIMIT ?
            """,
            (code, limit),
        ).fetchall()
    else:
        # Backward-compatible fallback until an existing database is re-imported.
        summary = conn.execute(
            f"""
            SELECT COUNT(*) AS products,
                   COUNT(DISTINCT NULLIF(manufacturer_inn,'')) AS manufacturers
            FROM gisp_products
            WHERE okpd2_code=? {where_active.replace('p.', '')}
            """,
            (code,),
        ).fetchone()
        rows = conn.execute(
            f"""
            SELECT registry_number, product_name, manufacturer_name, manufacturer_inn,
                   valid_until, tnved_code
            FROM gisp_products
            WHERE okpd2_code=? {where_active.replace('p.', '')}
            ORDER BY manufacturer_name, product_name, registry_number
            LIMIT ?
            """,
            (code, limit),
        ).fetchall()
    return {
        "okpd2_code": code,
        "active_only": active_only,
        "products": int(summary["products"] or 0),
        "manufacturers": int(summary["manufacturers"] or 0),
        "sample_products": [dict(row) for row in rows],
    }
